from io import BytesIO

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

from app.services.parsing import parse_document


def test_docx_parser_extracts_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_paragraph("AtlasAI enterprise knowledge")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "owner"
    table.rows[0].cells[1].text = "Soham"
    buffer = BytesIO()
    document.save(buffer)

    text = parse_document("brief.docx", buffer.getvalue())

    assert "AtlasAI enterprise knowledge" in text
    assert "owner | Soham" in text


def test_xlsx_parser_extracts_sheet_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["metric", "value"])
    sheet.append(["recall@5", 0.92])
    buffer = BytesIO()
    workbook.save(buffer)

    text = parse_document("metrics.xlsx", buffer.getvalue())

    assert "## Sheet: Metrics" in text
    assert "recall@5 | 0.92" in text


def test_pptx_parser_extracts_slide_text() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "AtlasAI"
    buffer = BytesIO()
    presentation.save(buffer)

    text = parse_document("pitch.pptx", buffer.getvalue())

    assert "## Slide 1" in text
    assert "AtlasAI" in text
