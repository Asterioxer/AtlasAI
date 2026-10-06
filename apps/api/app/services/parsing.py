import csv
import json
import re
from email import policy
from email.parser import BytesParser
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from pypdf import PdfReader

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    import yaml
except ImportError:
    yaml = None


class UnsupportedDocumentType(ValueError):
    pass


TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".xml", ".html", ".htm",
    ".css", ".js", ".jsx", ".ts", ".tsx", ".py", ".java", ".kt", ".go", ".rs",
    ".c", ".h", ".cpp", ".hpp", ".cs", ".php", ".rb", ".swift", ".scala",
    ".sql", ".sh", ".bash", ".zsh", ".ps1", ".bat", ".tex", ".dockerfile",
}


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _strip_rtf(text: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\\d* ?", " ", text)
    text = re.sub(r"[{}]", " ", text)
    return re.sub(r"\\par ?", "\n", text)


def _parse_zip_xml(content: bytes, extensions: tuple[str, ...]) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            parts = []
            for name in archive.namelist():
                if name.lower().endswith(extensions):
                    parts.append(_decode(archive.read(name)))
            if not parts:
                raise UnsupportedDocumentType("Archive contains no readable XML/text parts")
            return "\n\n".join(parts)
    except BadZipFile as exc:
        raise UnsupportedDocumentType("Invalid ZIP-based document") from exc


def parse_document(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix in TEXT_SUFFIXES:
        text = _decode(content)
        if suffix in {".html", ".htm", ".xml"} and BeautifulSoup:
            text = BeautifulSoup(
                text,
                "html.parser" if suffix in {".html", ".htm"} else "xml",
            ).get_text("\n", strip=True)
        elif suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            rows = csv.reader(StringIO(text), delimiter=delimiter)
            text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        elif suffix == ".json":
            try:
                text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                pass
        elif suffix in {".yaml", ".yml"} and yaml:
            try:
                text = yaml.safe_dump(yaml.safe_load(text), sort_keys=False, allow_unicode=True)
            except yaml.YAMLError:
                pass
        return text.strip()

    if suffix == ".rtf":
        return _strip_rtf(_decode(content)).strip()

    if suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(content)
        parts = []
        if message["subject"]:
            parts.append(f"Subject: {message['subject']}")
        for part in message.walk():
            if part.get_content_type() == "text/plain" and not part.get_content_disposition():
                try:
                    parts.append(part.get_content())
                except (LookupError, AttributeError):
                    parts.append(_decode(part.get_payload(decode=True) or b""))
        return "\n\n".join(parts).strip()

    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if suffix == ".docx":
        if DocxDocument is None:
            raise UnsupportedDocumentType("DOCX support is not installed")
        document = DocxDocument(BytesIO(content))
        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(blocks).strip()

    if suffix == ".xlsx":
        if load_workbook is None:
            raise UnsupportedDocumentType("XLSX support is not installed")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        blocks = []
        for sheet in workbook.worksheets:
            blocks.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    blocks.append(" | ".join(values))
        return "\n".join(blocks).strip()

    if suffix == ".xls":
        if xlrd is None:
            raise UnsupportedDocumentType("XLS support is not installed")
        workbook = xlrd.open_workbook(file_contents=content)
        blocks = []
        for sheet in workbook.sheets():
            blocks.append(f"## Sheet: {sheet.name}")
            for row_idx in range(sheet.nrows):
                values = [str(v) for v in sheet.row_values(row_idx)]
                if any(values):
                    blocks.append(" | ".join(values))
        return "\n".join(blocks).strip()

    if suffix == ".pptx":
        if Presentation is None:
            raise UnsupportedDocumentType("PPTX support is not installed")
        presentation = Presentation(BytesIO(content))
        blocks = []
        for index, slide in enumerate(presentation.slides, start=1):
            blocks.append(f"## Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    blocks.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(blocks).strip()

    if suffix in {".odt", ".ods", ".odp"}:
        return _parse_zip_xml(content, (".xml",)).strip()

    if suffix == ".epub":
        return _parse_zip_xml(content, (".html", ".xhtml", ".xml")).strip()

    raise UnsupportedDocumentType(f"Unsupported document type: {suffix or 'unknown'}")
